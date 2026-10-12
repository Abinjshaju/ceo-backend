"""Extract text from uploaded files while keeping pages, headings, sheets and table
headers, then chunk it into searchable sections."""

import csv
import io
import re
import zipfile
from dataclasses import dataclass
from pathlib import PurePosixPath
from xml.etree import ElementTree

EXTRACTOR_VERSION = "1"

SUPPORTED_TYPES = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".csv": "text/csv",
    ".txt": "text/plain",
    ".md": "text/markdown",
}

CHUNK_TARGET = 1500
CHUNK_OVERLAP = 200


class ExtractionError(Exception):
    """The file is unreadable, encrypted, or malformed."""


class NeedsOCR(Exception):
    """The PDF has no text layer."""


@dataclass
class RawSection:
    text: str
    heading: str | None = None
    page: int | None = None
    sheet: str | None = None
    table_header: str | None = None


@dataclass
class Chunk:
    text: str
    heading: str | None = None
    page: int | None = None
    sheet: str | None = None


def file_kind(filename: str) -> str | None:
    suffix = PurePosixPath(filename.lower()).suffix
    return suffix if suffix in SUPPORTED_TYPES else None


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ExtractionError("Could not decode the text file.")


def _pdf(data: bytes) -> tuple[list[RawSection], int]:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted and not reader.decrypt(""):
            raise ExtractionError("The PDF is password protected.")
        pages = [(page.extract_text() or "").strip() for page in reader.pages]
    except PdfReadError as exc:
        raise ExtractionError(f"The PDF could not be read: {exc}") from exc
    if sum(len(text) for text in pages) < 20 * max(len(pages), 1):
        raise NeedsOCR()
    return [RawSection(text=text, page=i + 1) for i, text in enumerate(pages) if text], len(pages)


W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _docx(data: bytes) -> tuple[list[RawSection], int | None]:
    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
        root = ElementTree.fromstring(archive.read("word/document.xml"))
    except (zipfile.BadZipFile, KeyError, ElementTree.ParseError) as exc:
        raise ExtractionError("The Word document could not be read.") from exc

    body = root.find(f"{W}body")
    sections: list[RawSection] = []
    heading: str | None = None
    buffer: list[str] = []

    def flush() -> None:
        text = "\n".join(buffer).strip()
        if text:
            sections.append(RawSection(text=text, heading=heading))
        buffer.clear()

    for element in list(body) if body is not None else []:
        if element.tag == f"{W}p":
            text = "".join(node.text or "" for node in element.iter(f"{W}t")).strip()
            style = element.find(f"{W}pPr/{W}pStyle")
            style_name = (style.get(f"{W}val") or "").lower() if style is not None else ""
            if text and (style_name.startswith("heading") or style_name == "title"):
                flush()
                heading = text
            elif text:
                buffer.append(text)
        elif element.tag == f"{W}tbl":
            flush()
            rows = []
            for row in element.iter(f"{W}tr"):
                cells = [
                    " ".join("".join(n.text or "" for n in cell.iter(f"{W}t")).split())
                    for cell in row.iter(f"{W}tc")
                ]
                if any(cells):
                    rows.append(" | ".join(cells))
            if rows:
                sections.append(
                    RawSection(text="\n".join(rows[1:]) or rows[0], heading=heading,
                               table_header=rows[0] if len(rows) > 1 else None)
                )
    flush()
    return sections, None


def _rows_to_section(rows: list[list[str]], sheet: str | None) -> RawSection | None:
    lines = [" | ".join(cell for cell in row) for row in rows if any(cell.strip() for cell in row)]
    if not lines:
        return None
    return RawSection(text="\n".join(lines[1:]) or lines[0], sheet=sheet,
                      table_header=lines[0] if len(lines) > 1 else None)


def _xlsx(data: bytes) -> tuple[list[RawSection], int | None]:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise ExtractionError("The spreadsheet could not be read.") from exc
    sections = []
    for sheet in workbook.worksheets:
        rows = [
            ["" if value is None else str(value).strip() for value in row]
            for row in sheet.iter_rows(values_only=True)
        ]
        section = _rows_to_section(rows, sheet.title)
        if section:
            sections.append(section)
    workbook.close()
    return sections, None


def _csv(data: bytes) -> tuple[list[RawSection], int | None]:
    rows = [[cell.strip() for cell in row] for row in csv.reader(io.StringIO(_decode(data)))]
    section = _rows_to_section(rows, None)
    return ([section] if section else []), None


def _markdown(data: bytes) -> tuple[list[RawSection], int | None]:
    sections: list[RawSection] = []
    heading: str | None = None
    buffer: list[str] = []
    for line in _decode(data).splitlines():
        match = re.match(r"^\s{0,3}#{1,6}\s+(.*)$", line)
        if match:
            if "".join(buffer).strip():
                sections.append(RawSection(text="\n".join(buffer).strip(), heading=heading))
            heading, buffer = match.group(1).strip(), []
        else:
            buffer.append(line)
    if "".join(buffer).strip():
        sections.append(RawSection(text="\n".join(buffer).strip(), heading=heading))
    return sections, None


def _text(data: bytes) -> tuple[list[RawSection], int | None]:
    text = _decode(data).strip()
    return ([RawSection(text=text)] if text else []), None


EXTRACTORS = {".pdf": _pdf, ".docx": _docx, ".xlsx": _xlsx, ".csv": _csv, ".md": _markdown, ".txt": _text}


def extract(data: bytes, filename: str) -> tuple[list[RawSection], int | None]:
    kind = file_kind(filename)
    if kind is None:
        raise ExtractionError("Unsupported file type.")
    return EXTRACTORS[kind](data)


def _split_long(text: str, target: int) -> list[str]:
    sentences = re.split(r"(?<=[.!?])\s+", text)
    pieces, current = [], ""
    for sentence in sentences:
        while len(sentence) > target:
            pieces.append(sentence[:target])
            sentence = sentence[target:]
        if len(current) + len(sentence) + 1 > target and current:
            pieces.append(current)
            current = sentence
        else:
            current = f"{current} {sentence}".strip()
    if current:
        pieces.append(current)
    return pieces


def _overlap_tail(text: str, size: int) -> str:
    tail = text[-size:]
    space = tail.find(" ")
    return tail[space + 1 :] if 0 <= space < len(tail) - 1 else tail


def chunk_sections(sections: list[RawSection], target: int = CHUNK_TARGET,
                   overlap: int = CHUNK_OVERLAP) -> list[Chunk]:
    chunks: list[Chunk] = []
    for section in sections:
        meta = {"heading": section.heading, "page": section.page, "sheet": section.sheet}
        if section.table_header:
            # Tables keep their header row in every chunk so values stay labelled.
            current: list[str] = []
            size = 0
            for row in section.text.splitlines():
                if current and size + len(row) > target:
                    chunks.append(Chunk(text="\n".join([section.table_header, *current]), **meta))
                    current, size = [], 0
                current.append(row)
                size += len(row) + 1
            if current:
                chunks.append(Chunk(text="\n".join([section.table_header, *current]), **meta))
            continue

        paragraphs = [p.strip() for p in re.split(r"\n\s*\n", section.text) if p.strip()]
        pieces: list[str] = []
        for paragraph in paragraphs:
            pieces.extend(_split_long(paragraph, target) if len(paragraph) > target else [paragraph])
        current_text = ""
        for piece in pieces:
            if current_text and len(current_text) + len(piece) + 2 > target:
                chunks.append(Chunk(text=current_text, **meta))
                current_text = f"{_overlap_tail(current_text, overlap)}\n\n{piece}"
            else:
                current_text = f"{current_text}\n\n{piece}".strip()
        if current_text:
            chunks.append(Chunk(text=current_text, **meta))
    return chunks
